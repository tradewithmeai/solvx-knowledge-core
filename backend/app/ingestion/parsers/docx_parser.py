"""DOCX file parser with heading and table support."""

import logging
from pathlib import Path
from typing import List, Optional

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_PARAGRAPH_STYLE
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from backend.app.ingestion.parsers.base import ParsedBlock, ParsedDocument, ParserBase
from backend.app.logging import get_logger
from backend.app.utils.errors import ParsingError

logger = get_logger(__name__)


class DocxParser(ParserBase):
    """Parser for DOCX (Microsoft Word) files."""

    SUPPORTED_EXTENSIONS = {".docx"}
    HEADING_STYLES = {
        "Heading 1",
        "Heading 2",
        "Heading 3",
        "Heading 4",
        "Heading 5",
        "Heading 6",
    }

    @classmethod
    def can_parse(cls, file_path: Path, mime_type: str | None = None) -> bool:
        """Check if file is a DOCX file."""
        if file_path.suffix.lower() in cls.SUPPORTED_EXTENSIONS:
            return True
        if mime_type and "wordprocessingml" in mime_type.lower():
            return True
        return False

    def parse(self, file_path: Path) -> ParsedDocument:
        """
        Parse DOCX file with heading structure and table support.

        Args:
            file_path: Path to DOCX file

        Returns:
            ParsedDocument with extracted text and structure

        Raises:
            ParsingError: If file cannot be parsed
        """
        try:
            logger.info(f"Parsing DOCX file: {file_path}")

            # Check file size
            file_size = file_path.stat().st_size
            if file_size == 0:
                logger.warning(f"Empty DOCX file: {file_path}")
                return ParsedDocument(blocks=[], metadata={"empty": True})

            # Load document with error handling
            try:
                doc = Document(file_path)
            except Exception as e:
                logger.error(f"Failed to load DOCX file {file_path}: {e}")
                raise ParsingError(
                    f"Failed to load DOCX file (corrupted or invalid): {e}",
                    {"file_path": str(file_path), "error": str(e)},
                )

            # Extract title from core properties
            title = self._extract_title(doc)

            # Parse document content
            blocks = self._parse_document_content(doc)

            # Build metadata
            metadata = {
                "file_size": file_size,
                "block_count": len(blocks),
                "paragraph_count": len([b for b in blocks if not b.section]),
                "section_count": len(set(b.section for b in blocks if b.section)),
            }

            # Add core properties metadata if available
            if doc.core_properties:
                try:
                    metadata["author"] = doc.core_properties.author or None
                    metadata["created"] = str(doc.core_properties.created) if doc.core_properties.created else None
                    metadata["modified"] = str(doc.core_properties.modified) if doc.core_properties.modified else None
                    metadata["subject"] = doc.core_properties.subject or None
                except Exception as e:
                    logger.debug(f"Could not extract core properties: {e}")

            logger.info(f"Successfully parsed DOCX file: {file_path} ({len(blocks)} blocks)")

            return ParsedDocument(
                blocks=blocks,
                title=title,
                metadata=metadata,
            )

        except ParsingError:
            raise
        except Exception as e:
            logger.error(f"Unexpected error parsing DOCX file {file_path}: {e}", exc_info=True)
            raise ParsingError(
                f"Unexpected error while parsing DOCX file: {e}",
                {"file_path": str(file_path), "error": str(e)},
            )

    def _extract_title(self, doc: DocxDocument) -> Optional[str]:
        """
        Extract document title from core properties or first heading.

        Args:
            doc: Loaded DOCX document

        Returns:
            Document title or None
        """
        try:
            # Try to get title from core properties
            if doc.core_properties and doc.core_properties.title:
                return doc.core_properties.title.strip()
        except Exception as e:
            logger.debug(f"Could not extract title from core properties: {e}")

        # Try to get title from first heading
        try:
            for paragraph in doc.paragraphs:
                if self._is_heading(paragraph):
                    text = paragraph.text.strip()
                    if text:
                        return text
        except Exception as e:
            logger.debug(f"Could not extract title from headings: {e}")

        return None

    def _parse_document_content(self, doc: DocxDocument) -> List[ParsedBlock]:
        """
        Parse all content from DOCX document.

        Args:
            doc: Loaded DOCX document

        Returns:
            List of ParsedBlocks
        """
        blocks = []
        current_section = None
        current_text_parts = []
        byte_offset = 0

        try:
            for element in doc.element.body:
                try:
                    # Handle paragraphs
                    if element.tag.endswith("p"):
                        # Get paragraph object
                        para = self._get_paragraph_from_element(element, doc)
                        if para is None:
                            continue

                        # Check if this is a heading
                        if self._is_heading(para):
                            # Save previous section if exists
                            if current_text_parts:
                                text = "\n".join(current_text_parts).strip()
                                if text:
                                    block = self._create_block(
                                        text=text,
                                        section=current_section,
                                        offset_start=byte_offset - len("\n".join(current_text_parts).encode("utf-8")),
                                        offset_end=byte_offset,
                                    )
                                    blocks.append(block)
                                current_text_parts = []

                            # Update current section
                            current_section = para.text
                        else:
                            # Add paragraph text to current section
                            text = para.text.strip()
                            if text:
                                current_text_parts.append(text)
                                byte_offset += len(text.encode("utf-8")) + 1

                    # Handle tables
                    elif element.tag.endswith("tbl"):
                        # Save previous section if exists
                        if current_text_parts:
                            text = "\n".join(current_text_parts).strip()
                            if text:
                                block = self._create_block(
                                    text=text,
                                    section=current_section,
                                    offset_start=byte_offset - len("\n".join(current_text_parts).encode("utf-8")),
                                    offset_end=byte_offset,
                                    block_type="paragraph",
                                )
                                blocks.append(block)
                            current_text_parts = []

                        # Convert table to text
                        table = self._get_table_from_element(element, doc)
                        if table is not None:
                            table_text = self._convert_table_to_text(table)
                            if table_text.strip():
                                block = self._create_block(
                                    text=table_text,
                                    section=current_section,
                                    offset_start=byte_offset,
                                    offset_end=byte_offset + len(table_text.encode("utf-8")),
                                    block_type="table",
                                )
                                blocks.append(block)
                                byte_offset += len(table_text.encode("utf-8")) + 1

                except Exception as e:
                    logger.warning(f"Error processing element: {e}")
                    continue

        except Exception as e:
            logger.warning(f"Error iterating document elements: {e}")

        # Add final section if exists
        if current_text_parts:
            text = "\n".join(current_text_parts).strip()
            if text:
                block = self._create_block(
                    text=text,
                    section=current_section,
                    offset_start=byte_offset - len(text.encode("utf-8")),
                    offset_end=byte_offset,
                    block_type="paragraph",
                )
                blocks.append(block)

        # If no blocks were extracted, try fallback parsing
        if not blocks:
            blocks = self._parse_paragraphs_fallback(doc)

        return blocks

    def _parse_paragraphs_fallback(self, doc: DocxDocument) -> List[ParsedBlock]:
        """
        Fallback parsing using doc.paragraphs when element iteration fails.

        Args:
            doc: Loaded DOCX document

        Returns:
            List of ParsedBlocks
        """
        blocks = []
        current_section = None
        current_text_parts = []
        byte_offset = 0

        try:
            for para in doc.paragraphs:
                try:
                    if self._is_heading(para):
                        # Save previous section
                        if current_text_parts:
                            text = "\n".join(current_text_parts).strip()
                            if text:
                                blocks.append(
                                    self._create_block(
                                        text=text,
                                        section=current_section,
                                        offset_start=byte_offset - len(text.encode("utf-8")),
                                        offset_end=byte_offset,
                                    )
                                )
                            current_text_parts = []

                        # Update section
                        current_section = para.text
                    else:
                        # Add paragraph text
                        text = para.text.strip()
                        if text:
                            current_text_parts.append(text)
                            byte_offset += len(text.encode("utf-8")) + 1

                except Exception as e:
                    logger.debug(f"Error processing paragraph: {e}")
                    continue

            # Add final section
            if current_text_parts:
                text = "\n".join(current_text_parts).strip()
                if text:
                    blocks.append(
                        self._create_block(
                            text=text,
                            section=current_section,
                            offset_start=byte_offset - len(text.encode("utf-8")),
                            offset_end=byte_offset,
                        )
                    )

        except Exception as e:
            logger.error(f"Fallback paragraph parsing failed: {e}")

        return blocks

    def _is_heading(self, para: Paragraph) -> bool:
        """
        Check if paragraph is a heading.

        Args:
            para: Paragraph to check

        Returns:
            True if paragraph is a heading style
        """
        try:
            if para.style and para.style.name in self.HEADING_STYLES:
                return True

            # Also check for direct style property
            pPr = para._element.get_or_add_pPr()
            pStyle = pPr.find(qn("w:pStyle"))
            if pStyle is not None:
                style_name = pStyle.get(qn("w:val"))
                if style_name and style_name.startswith("Heading"):
                    return True

        except Exception as e:
            logger.debug(f"Error checking heading style: {e}")

        return False

    def _convert_table_to_text(self, table: Table) -> str:
        """
        Convert a Word table to a plain text representation.

        Args:
            table: DOCX table object

        Returns:
            Text representation of the table
        """
        try:
            lines = []
            for row_idx, row in enumerate(table.rows):
                try:
                    cells = []
                    for cell in row.cells:
                        try:
                            cell_text = self._extract_cell_text(cell)
                            cells.append(cell_text)
                        except Exception as e:
                            logger.debug(f"Error extracting cell text: {e}")
                            cells.append("")

                    # Join cells with pipe separator
                    line = " | ".join(cells)
                    lines.append(line)

                    # Add separator after header row (first row)
                    if row_idx == 0:
                        lines.append("-" * (len(line) + 10))

                except Exception as e:
                    logger.debug(f"Error processing table row: {e}")
                    continue

            return "\n".join(lines)

        except Exception as e:
            logger.error(f"Error converting table to text: {e}")
            return ""

    def _extract_cell_text(self, cell: _Cell) -> str:
        """
        Extract text content from a table cell.

        Args:
            cell: Table cell object

        Returns:
            Text content of the cell
        """
        try:
            texts = []
            for para in cell.paragraphs:
                try:
                    text = para.text.strip()
                    if text:
                        texts.append(text)
                except Exception as e:
                    logger.debug(f"Error extracting paragraph from cell: {e}")

            return " ".join(texts) if texts else ""

        except Exception as e:
            logger.debug(f"Error extracting cell text: {e}")
            return ""

    def _get_paragraph_from_element(self, element, doc: DocxDocument) -> Optional[Paragraph]:
        """
        Get Paragraph object from an XML element.

        Args:
            element: XML element
            doc: DOCX document

        Returns:
            Paragraph object or None
        """
        try:
            for para in doc.paragraphs:
                if para._element == element:
                    return para
        except Exception as e:
            logger.debug(f"Error getting paragraph from element: {e}")

        return None

    def _get_table_from_element(self, element, doc: DocxDocument) -> Optional[Table]:
        """
        Get Table object from an XML element.

        Args:
            element: XML element
            doc: DOCX document

        Returns:
            Table object or None
        """
        try:
            for table in doc.tables:
                if table._element == element:
                    return table
        except Exception as e:
            logger.debug(f"Error getting table from element: {e}")

        return None
